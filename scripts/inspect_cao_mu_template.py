from pathlib import Path
from openpyxl import load_workbook

path = Path('/home/ubuntu/upload/coord_signal_response/VqqcDNFbm9mRHLM4Zi2Pw8/Tong_hop_so_lieu_hang_ngay_2026.xlsx')
wb = load_workbook(path, data_only=False, read_only=True)
print('SHEETS', wb.sheetnames)
for ws in wb.worksheets:
    print(f'--- SHEET {ws.title} rows={ws.max_row} cols={ws.max_column}')
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 12), values_only=False):
        cells = []
        for cell in row:
            if cell.value is not None:
                cells.append(f'{cell.coordinate}={cell.value!r}<{cell.data_type}>')
        if cells:
            print(' | '.join(cells))
