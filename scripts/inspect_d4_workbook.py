from pathlib import Path
from openpyxl import load_workbook

path = Path('/home/ubuntu/upload/coord_signal_response/1P2tPKx3fIEQ6aViZnJNo2/mau-import-PhanchianhâncongvươncayD4.xlsx')
wb = load_workbook(path, data_only=False, read_only=True)
print('SHEETS', wb.sheetnames)
for ws in wb.worksheets:
    print(f'\nSHEET {ws.title!r} rows={ws.max_row} cols={ws.max_column}')
    for row_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=ws.max_row, values_only=False), start=1):
        values = []
        for col_idx, cell in enumerate(row, start=1):
            value = cell.value
            if value is not None:
                values.append({'coord': cell.coordinate if hasattr(cell, 'coordinate') else f'{col_idx}{row_idx}', 'value': value, 'type': getattr(cell, 'data_type', None), 'number_format': getattr(cell, 'number_format', None)})
        if row_idx <= 5 or row_idx >= 25 or any(item['coord'][0] in 'FJNP' and isinstance(item['value'], (int, float)) and not float(item['value']).is_integer() for item in values):
            print(row_idx, values)
