from pathlib import Path
from openpyxl import load_workbook

path = Path('/home/ubuntu/upload/coord_signal_response/1P2tPKx3fIEQ6aViZnJNo2/mau-import-PhanchianhâncongvươncayD4.xlsx')
wb = load_workbook(path, data_only=False, read_only=True)
ws = wb['Phân chia nhân công vườn cây']
for row_idx in range(3, ws.max_row + 1):
    values = []
    for col in ('A', 'B', 'F', 'J', 'N', 'O', 'P'):
        cell = ws[f'{col}{row_idx}']
        values.append((col, cell.value, cell.data_type, cell.number_format))
    print(row_idx, values)
print('NON_INTEGER_TREE_CELLS')
for row_idx in range(3, ws.max_row + 1):
    for col in ('F', 'J', 'N'):
        value = ws[f'{col}{row_idx}'].value
        if isinstance(value, (int, float)) and not float(value).is_integer():
            print(row_idx, col, repr(value), ws[f'{col}{row_idx}'].data_type)
